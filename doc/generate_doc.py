"""
Generates a simple, easy-to-read .docx explaining Legilimens — with a mind map,
an architecture diagram, and a separate diagram for EACH step of the proxy flow.

Self-contained: everything (this script, its deps, the diagrams, the output) lives in
/doc and never touches the main codebase.

Run:
    .venv\\Scripts\\python.exe doc\\generate_doc.py      # Windows
    .venv/bin/python doc/generate_doc.py                # macOS/Linux

Output (gitignored):
    doc/assets/*.png        intermediate diagram images
    doc/Legilimens.docx     the document
"""

from pathlib import Path

import matplotlib
matplotlib.use("Agg")  # no display needed
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---------- palette (clean, light, readable) ----------
ACCENT = "#C8F400"   # highlight (brand acid-green)
DARK = "#1a1a1a"
GREY = "#c2c2c2"
MUTED = "#8a8a8a"
WHITE = "#ffffff"

DOC_DIR = Path(__file__).parent
ASSETS = DOC_DIR / "assets"
ASSETS.mkdir(exist_ok=True)

# ---------- shared box layout for architecture + step diagrams ----------
BOXES = {
    "client": {"xy": (0.4, 3.3), "w": 2.4, "h": 1.1, "label": "Game / App\nClient"},
    "proxy":  {"xy": (3.8, 3.3), "w": 2.4, "h": 1.1, "label": "LEGILIMENS\nProxy  :4433"},
    "server": {"xy": (7.2, 3.3), "w": 2.4, "h": 1.1, "label": "Real Server\n:4434"},
    "ui":     {"xy": (3.8, 0.7), "w": 2.4, "h": 1.1, "label": "Inspector UI\n(React)"},
}

ARROWS = {
    "c2p": (("client", "r"), ("proxy", "l")),
    "p2s": (("proxy", "r"), ("server", "l")),
    "p2u": (("proxy", "b"), ("ui", "t")),
}


def _edges(b):
    x, y = b["xy"]; w, h = b["w"], b["h"]
    return {
        "c": (x + w / 2, y + h / 2),
        "l": (x, y + h / 2),
        "r": (x + w, y + h / 2),
        "t": (x + w / 2, y + h),
        "b": (x + w / 2, y),
    }


def draw_scene(active_boxes, active_arrows, title, note=None, fname="scene.png",
               dim_inactive=True):
    """One frame of the pipeline. active_arrows = {key: label}.

    dim_inactive=True  -> inactive boxes fade (used for the per-step diagrams).
    dim_inactive=False -> all boxes stay readable (used for the architecture overview).
    """
    fig, ax = plt.subplots(figsize=(9, 5.2))
    ax.set_xlim(0, 10); ax.set_ylim(0, 5.9); ax.axis("off")

    # arrows (under boxes); labels float ABOVE the boxes (horizontal) or beside (vertical)
    for key, ((b1, e1), (b2, e2)) in ARROWS.items():
        p1 = _edges(BOXES[b1])[e1]
        p2 = _edges(BOXES[b2])[e2]
        on = key in active_arrows
        ax.annotate("", xy=p2, xytext=p1,
                    arrowprops=dict(arrowstyle="-|>", lw=3.2 if on else 1.1,
                                    color=DARK if on else GREY, shrinkA=3, shrinkB=3))
        if on and active_arrows[key]:
            if key == "p2u":  # vertical arrow -> label to the right, in open space
                ax.text(p1[0] + 0.9, (p1[1] + p2[1]) / 2, active_arrows[key],
                        ha="left", va="center", fontsize=10.5, color=DARK, fontweight="bold")
            else:             # horizontal arrow -> label above the boxes (no overlap)
                mx = (p1[0] + p2[0]) / 2
                ax.text(mx, 4.78, active_arrows[key], ha="center", va="center",
                        fontsize=10.5, color=DARK, fontweight="bold")

    # boxes
    for name, b in BOXES.items():
        on = name in active_boxes
        if on:
            fc, ec, tc, lw, fw = ACCENT, DARK, DARK, 2.6, "bold"
        elif dim_inactive:
            fc, ec, tc, lw, fw = WHITE, GREY, MUTED, 1.3, "normal"
        else:
            fc, ec, tc, lw, fw = WHITE, DARK, DARK, 1.6, "normal"
        ax.add_patch(Rectangle(b["xy"], b["w"], b["h"], linewidth=lw,
                               edgecolor=ec, facecolor=fc, zorder=2))
        cx, cy = _edges(b)["c"]
        ax.text(cx, cy, b["label"], ha="center", va="center", zorder=3,
                fontsize=11, color=tc, fontweight=fw)

    # optional note above the proxy (for the read / tamper steps that have no arrow)
    if note:
        px = _edges(BOXES["proxy"])["t"][0]
        ax.text(px, 4.8, note, ha="center", va="center", fontsize=11,
                color=DARK, fontweight="bold",
                bbox=dict(boxstyle="round,pad=0.3", facecolor=ACCENT, edgecolor=DARK, lw=1.4))

    ax.text(5, 5.6, title, ha="center", va="center", fontsize=13.5,
            fontweight="bold", color=DARK)

    fig.savefig(ASSETS / fname, dpi=150, bbox_inches="tight")
    plt.close(fig)


def draw_architecture():
    # Calm overview: every box readable, only the proxy highlighted (it's the star).
    draw_scene(
        active_boxes={"proxy"},
        active_arrows={"c2p": "WebTransport", "p2s": "forwards", "p2u": "live copy (WS)"},
        title="Architecture — the whole picture",
        fname="architecture.png",
        dim_inactive=False,
    )


def draw_frontend():
    """Frontend architecture: one-way data IN, controls OUT, and the loop between them."""
    fig, ax = plt.subplots(figsize=(9.5, 5.3))
    ax.set_xlim(0, 12); ax.set_ylim(0, 6); ax.axis("off")

    def box(x, y, w, h, label, hot=False, fs=10.5):
        ax.add_patch(Rectangle((x, y), w, h, linewidth=2.6 if hot else 1.6,
                               edgecolor=DARK, facecolor=ACCENT if hot else WHITE, zorder=2))
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", zorder=3,
                fontsize=fs, color=DARK, fontweight="bold" if hot else "normal")

    def arrow(p1, p2, label="", dy=0.3, color=DARK, lw=2.6, style="-", fs=9.5, italic=False):
        ax.annotate("", xy=p2, xytext=p1,
                    arrowprops=dict(arrowstyle="-|>", lw=lw, color=color, shrinkA=3, shrinkB=3,
                                    linestyle=style))
        if label:
            mx, my = (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2
            ax.text(mx, my + dy, label, ha="center", va="center", fontsize=fs,
                    color=color, fontweight="bold" if not italic else "normal",
                    style="italic" if italic else "normal")

    # Top row — DATA IN (read path)
    box(0.3, 3.9, 2.3, 1.1, "WebSocket\n:4435")
    box(3.8, 3.9, 3.3, 1.1, "Zustand Store\n(single source of truth)", hot=True, fs=9.5)
    box(8.2, 3.5, 3.6, 1.9, "View Components\nTrafficLog · LatencyGraph\nStatusBar · StreamInspector")
    arrow((2.6, 4.45), (3.8, 4.45), "events")
    arrow((7.1, 4.45), (8.2, 4.45), "renders")

    # Bottom row — CONTROLS OUT (write path)
    box(0.3, 0.8, 2.3, 1.1, "Backend\nControl API\n:4436", fs=9.5)
    box(4.0, 0.7, 4.5, 1.3, "Control Components\nHeader · TargetConfig\nTamperConfig · AttackSimulator")
    arrow((4.0, 1.35), (2.6, 1.35), "REST POST")  # controls -> API (points left)

    # Loop close — API back up to WebSocket (the backend reacts, new events flow in)
    arrow((1.45, 1.9), (1.45, 3.9), color=MUTED, lw=1.4, style=(0, (4, 3)))
    ax.text(1.75, 2.9, "backend reacts →\nnew events\nflow back in", ha="left", va="center",
            fontsize=8.5, color=MUTED, style="italic")

    ax.text(6, 5.75, "Frontend Architecture — data in, controls out",
            ha="center", va="center", fontsize=13.5, fontweight="bold", color=DARK)
    fig.savefig(ASSETS / "frontend.png", dpi=150, bbox_inches="tight")
    plt.close(fig)


# ---------- frontend STEP diagrams (same layout as the overview above) ----------
FE_BOXES = {
    "ws":       {"xy": (0.3, 3.9), "w": 2.3, "h": 1.1, "label": "WebSocket\n:4435"},
    "store":    {"xy": (3.8, 3.9), "w": 3.3, "h": 1.1, "label": "Zustand Store\n(single source of truth)", "fs": 9.5},
    "views":    {"xy": (8.2, 3.5), "w": 3.6, "h": 1.9, "label": "View Components\nTrafficLog · LatencyGraph\nStatusBar · StreamInspector"},
    "api":      {"xy": (0.3, 0.8), "w": 2.3, "h": 1.1, "label": "Backend\nControl API\n:4436", "fs": 9.5},
    "controls": {"xy": (4.0, 0.7), "w": 4.5, "h": 1.3, "label": "Control Components\nHeader · TargetConfig\nTamperConfig · AttackSimulator"},
}

FE_ARROWS = {
    "ws2store":     (("ws", "r"), ("store", "l")),
    "store2views":  (("store", "r"), ("views", "l")),
    "controls2api": (("controls", "l"), ("api", "r")),  # points left
}


def draw_fe(active_boxes, active_arrows, title, note=None, fname="fe.png",
            show_loop=False, loop_active=False):
    """One frame of the frontend pipeline. note=(text, box_key) draws a badge by a box."""
    fig, ax = plt.subplots(figsize=(9.5, 5.4))
    ax.set_xlim(0, 12); ax.set_ylim(0, 6.1); ax.axis("off")

    for key, ((b1, e1), (b2, e2)) in FE_ARROWS.items():
        p1 = _edges(FE_BOXES[b1])[e1]; p2 = _edges(FE_BOXES[b2])[e2]
        on = key in active_arrows
        ax.annotate("", xy=p2, xytext=p1,
                    arrowprops=dict(arrowstyle="-|>", lw=3.0 if on else 1.1,
                                    color=DARK if on else GREY, shrinkA=3, shrinkB=3))
        if on and active_arrows[key]:
            mx, my = (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2
            ax.text(mx, my + 0.3, active_arrows[key], ha="center", va="center",
                    fontsize=9.8, color=DARK, fontweight="bold")

    if show_loop:
        p1 = _edges(FE_BOXES["api"])["t"]; p2 = _edges(FE_BOXES["ws"])["b"]
        col = DARK if loop_active else MUTED
        ax.annotate("", xy=p2, xytext=p1,
                    arrowprops=dict(arrowstyle="-|>", lw=2.2 if loop_active else 1.3,
                                    color=col, shrinkA=3, shrinkB=3, linestyle=(0, (4, 3))))
        ax.text(1.75, 2.9, "backend reacts,\nnew events\nflow back in", ha="left",
                va="center", fontsize=8.5, color=col, style="italic")

    for name, b in FE_BOXES.items():
        on = name in active_boxes
        if on:
            fc, ec, tc, lw, fw = ACCENT, DARK, DARK, 2.6, "bold"
        else:
            fc, ec, tc, lw, fw = WHITE, GREY, MUTED, 1.3, "normal"
        ax.add_patch(Rectangle(b["xy"], b["w"], b["h"], linewidth=lw,
                               edgecolor=ec, facecolor=fc, zorder=2))
        cx, cy = _edges(b)["c"]
        ax.text(cx, cy, b["label"], ha="center", va="center", zorder=3,
                fontsize=b.get("fs", 10.5), color=tc, fontweight=fw)

    if note:
        text, anchor = note
        b = FE_BOXES[anchor]
        cx = _edges(b)["c"][0]
        top = b["xy"][1] + b["h"]
        by = top + 0.34 if top + 0.5 <= 5.6 else b["xy"][1] - 0.34
        ax.text(cx, by, text, ha="center", va="center", fontsize=10, color=DARK,
                fontweight="bold",
                bbox=dict(boxstyle="round,pad=0.3", facecolor=ACCENT, edgecolor=DARK, lw=1.4))

    ax.text(6, 5.9, title, ha="center", va="center", fontsize=13.5,
            fontweight="bold", color=DARK)
    fig.savefig(ASSETS / fname, dpi=150, bbox_inches="tight")
    plt.close(fig)


FE_STEPS = [
    (dict(active_boxes={"ws"}, active_arrows={}, title="Step 1 - Connect",
          note=("OPEN + AUTO-RECONNECT", "ws"), fname="fe1.png"),
     "When the page loads, the app opens a WebSocket to the backend. If it ever drops, it reconnects on its own."),

    (dict(active_boxes={"ws", "store"}, active_arrows={"ws2store": "by type"},
          title="Step 2 - Receive & Route", fname="fe2.png"),
     "Each event arrives as JSON on the socket. The store routes it by type: traffic, attack, or intercept."),

    (dict(active_boxes={"store"}, active_arrows={}, title="Step 3 - Update the Store",
          note=("ADD · COUNT · FLAG", "store"), fname="fe3.png"),
     "The store saves the event, updates the live counters, and flags risky payloads like a leaked token."),

    (dict(active_boxes={"store", "views"}, active_arrows={"store2views": "re-render"},
          title="Step 4 - Render", fname="fe4.png"),
     "Components that subscribe to the store re-render - and only the ones that use the changed data."),

    (dict(active_boxes={"views"}, active_arrows={}, title="Step 5 - Show It",
          note=("LIVE LOG · FILTER · SEARCH", "views"), fname="fe5.png"),
     "A new row appears in the live traffic log. Filters and search let you narrow down what you see."),

    (dict(active_boxes={"controls", "api"}, active_arrows={"controls2api": "REST POST"},
          title="Step 6 - Send a Control", fname="fe6.png"),
     "When you click START, PAUSE, or set a tamper rule, a control component sends a REST request to the backend."),

    (dict(active_boxes={"ws", "api"}, active_arrows={}, title="Step 7 - The Loop Closes",
          fname="fe7.png", show_loop=True, loop_active=True),
     "The backend reacts and sends new events back over the WebSocket - so your own actions show up in the log too."),
]


def draw_mindmap():
    fig, ax = plt.subplots(figsize=(9, 5.4))
    ax.set_xlim(0, 10); ax.set_ylim(0, 5.6); ax.axis("off")
    center = (5, 2.8)

    branches = [
        (1.7, 4.6, "PROBLEM", "WebTransport is invisible\nto Burp / Wireshark"),
        (1.7, 2.8, "WHAT", "A man-in-the-middle\nproxy + live inspector"),
        (1.7, 1.0, "TARGET", "Vault713 —\na vulnerable game"),
        (8.3, 4.6, "FRONTEND", "React · Zustand\nlive traffic log"),
        (8.3, 2.8, "BACKEND", "Python · aioquic\nFastAPI · WebSocket"),
        (8.3, 1.0, "YOU CAN", "read · tamper\nattack"),
    ]

    for bx, by, _, _ in branches:
        ax.plot([center[0], bx], [center[1], by], color=GREY, lw=1.4, zorder=1)

    # center
    cw, ch = 2.2, 1.1
    ax.add_patch(Rectangle((center[0] - cw / 2, center[1] - ch / 2), cw, ch,
                           linewidth=2.6, edgecolor=DARK, facecolor=ACCENT, zorder=3))
    ax.text(center[0], center[1], "LEGILIMENS", ha="center", va="center",
            fontsize=13, fontweight="bold", color=DARK, zorder=4)

    # branches
    for bx, by, t, d in branches:
        w, h = 2.7, 1.05
        ax.add_patch(Rectangle((bx - w / 2, by - h / 2), w, h,
                               linewidth=1.5, edgecolor=DARK, facecolor=WHITE, zorder=3))
        ax.text(bx, by + 0.2, t, ha="center", va="center", fontsize=10.5,
                fontweight="bold", color=DARK, zorder=4)
        ax.text(bx, by - 0.22, d, ha="center", va="center", fontsize=8.5,
                color="#555555", zorder=4)

    ax.text(5, 5.4, "Mind Map", ha="center", va="center", fontsize=13.5,
            fontweight="bold", color=DARK)
    fig.savefig(ASSETS / "mindmap.png", dpi=150, bbox_inches="tight")
    plt.close(fig)


# ---------- the 7 steps (one diagram each) ----------
STEPS = [
    (dict(active_boxes={"client", "proxy"}, active_arrows={"c2p": "sends a message"},
          title="Step 1 — Connect & Send", note=None, fname="step1.png"),
     "The app connects to Legilimens, thinking it is the real server, and sends a message."),

    (dict(active_boxes={"proxy", "server"}, active_arrows={"p2s": "connects"},
          title="Step 2 — Open the Real Server", note=None, fname="step2.png"),
     "Legilimens quietly opens its own connection to the real server behind the scenes."),

    (dict(active_boxes={"proxy"}, active_arrows={}, title="Step 3 — Read",
          note="READ & FLAG", fname="step3.png"),
     "Legilimens reads the message and flags anything risky, like a leaked token or password."),

    (dict(active_boxes={"proxy"}, active_arrows={}, title="Step 4 — Tamper (optional)",
          note="REWRITE FIELD", fname="step4.png"),
     "If a rule is on, Legilimens changes a value in the message, for example score becomes 99999."),

    (dict(active_boxes={"proxy", "server"}, active_arrows={"p2s": "forwards"},
          title="Step 5 — Forward", note=None, fname="step5.png"),
     "The message (changed or not) is passed on to the real server, which never notices."),

    (dict(active_boxes={"proxy", "ui"}, active_arrows={"p2u": "WebSocket"},
          title="Step 6 — Broadcast a Copy", note=None, fname="step6.png"),
     "At the same time, a copy of every message is pushed to the inspector UI over a WebSocket."),

    (dict(active_boxes={"ui"}, active_arrows={}, title="Step 7 — Show It Live",
          note=None, fname="step7.png"),
     "The UI shows the message instantly in a live log you can filter and search."),
]


def build_diagrams():
    draw_mindmap()
    draw_architecture()
    draw_frontend()
    for scene, _ in FE_STEPS:
        draw_fe(**scene)
    for scene, _ in STEPS:
        draw_scene(**scene)


# ---------- assemble the .docx ----------
def _heading(doc, text):
    h = doc.add_heading(text, level=1)
    return h


def build_docx():
    doc = Document()

    # Title
    t = doc.add_heading("Legilimens", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run('WebTransport Traffic Inspector — "Reading what others cannot see."')
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # In one line
    _heading(doc, "In one line")
    doc.add_paragraph(
        "Legilimens is a wiretap for WebTransport. It sits between an app and its server, "
        "shows you every message live, and lets you change or block them. Think of it as "
        "Burp Suite, but for traffic that normal tools cannot see."
    )

    # The problem
    _heading(doc, "The problem it solves")
    doc.add_paragraph(
        "WebTransport runs over QUIC and UDP, and it is encrypted. Tools like Burp Suite and "
        "Wireshark cannot read it. So when a developer builds a real-time app on WebTransport "
        "(a game, a live dashboard, a chat), there is no easy way to inspect or test its "
        "traffic. Legilimens fills that gap."
    )

    # Mind map
    _heading(doc, "Mind map")
    doc.add_picture(str(ASSETS / "mindmap.png"), width=Inches(6.3))
    _center_last_image(doc)

    # Architecture
    doc.add_page_break()
    _heading(doc, "Architecture")
    doc.add_paragraph(
        "The app connects to Legilimens instead of the real server. Legilimens forwards "
        "everything on, and sends a live copy of each message to the inspector UI."
    )
    doc.add_picture(str(ASSETS / "architecture.png"), width=Inches(6.3))
    _center_last_image(doc)

    # Frontend architecture (the part that matters for a frontend role)
    doc.add_page_break()
    _heading(doc, "Frontend architecture")
    doc.add_paragraph(
        "The inspector UI is a React app built on one rule: data flows one way. There are "
        "two paths, and together they form a loop."
    )
    doc.add_picture(str(ASSETS / "frontend.png"), width=Inches(6.5))
    _center_last_image(doc)

    p = doc.add_paragraph()
    p.add_run("Data in (read path). ").bold = True
    p.add_run(
        "A WebSocket pushes live events into a single Zustand store — the one source of "
        "truth. View components (the traffic log, charts, status bar, stream inspector) "
        "subscribe to the store and re-render. No component keeps its own copy of the data."
    )
    p = doc.add_paragraph()
    p.add_run("Controls out (write path). ").bold = True
    p.add_run(
        "When you click a button — start, pause, set a tamper rule, change the target, run "
        "an attack — a control component sends a REST request to the backend API. The backend "
        "acts, and the results come back as new events over the WebSocket. That closes the loop."
    )
    p = doc.add_paragraph()
    p.add_run("Why it's built this way. ").bold = True
    p.add_run(
        "Components are cleanly split into controls (which write) and views (which read), so "
        "the app is easy to reason about. The store stays pure render-state; the live socket "
        "and connection objects are kept outside it. And because the log can update tens of "
        "times a second, store selectors keep re-renders scoped to only the components that "
        "use each slice of state."
    )

    # Frontend — step by step
    doc.add_page_break()
    _heading(doc, "How the frontend works — step by step")
    doc.add_paragraph(
        "Follow one event from the wire to the screen, then how a click flows back out. "
        "Each step has its own picture; the highlighted boxes show what is active."
    )
    for scene, sentence in FE_STEPS:
        p = doc.add_paragraph()
        run = p.add_run(scene["title"])
        run.bold = True
        run.font.size = Pt(12)
        doc.add_paragraph(sentence)
        doc.add_picture(str(ASSETS / scene["fname"]), width=Inches(6.0))
        _center_last_image(doc)

    # Proxy — step by step
    doc.add_page_break()
    _heading(doc, "How the proxy works — step by step")
    doc.add_paragraph(
        "Now the backend side. Follow one message on its journey through the proxy. Each step "
        "is shown on its own picture; the highlighted boxes show what is active."
    )
    for scene, sentence in STEPS:
        p = doc.add_paragraph()
        run = p.add_run(scene["title"])
        run.bold = True
        run.font.size = Pt(12)
        doc.add_paragraph(sentence)
        doc.add_picture(str(ASSETS / scene["fname"]), width=Inches(5.7))
        _center_last_image(doc)

    # What you can do
    doc.add_page_break()
    _heading(doc, "What you can do with it")
    for line in [
        "Read — watch every message live and spot leaked secrets.",
        "Tamper — change a value in flight (e.g. force a score) to test if the server trusts it.",
        "Pause / Disconnect — hold or cut the traffic on demand.",
        "Attack — run connection floods and other stress tests against the server.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    # Tech stack
    _heading(doc, "Tech stack")
    doc.add_paragraph(
        "Frontend: React, TypeScript, Zustand, Recharts.    "
        "Backend: Python, aioquic (QUIC/WebTransport), FastAPI, WebSockets."
    )

    # Status
    _heading(doc, "Where it stands")
    doc.add_paragraph(
        "The core works end to end: a real man-in-the-middle proxy over QUIC, with a live "
        "React inspector on top. It is a strong prototype. Next steps include an interactive "
        "edit-and-forward panel, saving sessions, and finishing a couple of the attacks."
    )

    out = DOC_DIR / "Legilimens.docx"
    doc.save(str(out))
    return out


def _center_last_image(doc):
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


if __name__ == "__main__":
    print("[doc] building diagrams...")
    build_diagrams()
    print("[doc] assembling .docx...")
    out = build_docx()
    print(f"[doc] done -> {out}")
